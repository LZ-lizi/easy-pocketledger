"""Diff the edited copy workbook against the one the extractor produced."""

import os
import subprocess
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HERE)
ORIGINAL = os.path.join(os.environ.get("TEMP", "/tmp"), "orig.xlsx")

subprocess.run(["git", "-C", PROJECT, "show", "HEAD:tools/记账本文案.xlsx"],
               stdout=open(ORIGINAL, "wb"), check=True)

a = openpyxl.load_workbook(ORIGINAL)
b = openpyxl.load_workbook(os.path.join(HERE, "记账本文案.xlsx"))

total = 0
for name in a.sheetnames:
    sa, sb = a[name], b[name]
    rows = max(sa.max_row, sb.max_row)
    cols = max(sa.max_column, sb.max_column)
    diffs = []
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            va = sa.cell(row=r, column=c).value
            vb = sb.cell(row=r, column=c).value
            if va != vb:
                diffs.append((r, c, va, vb))
    total += len(diffs)
    print(f"===== {name}: {len(diffs)} changed =====")
    for r, c, va, vb in diffs:
        label = sa.cell(row=r, column=1).value
        where = sa.cell(row=r, column=2).value
        element = sa.cell(row=r, column=3).value
        scene = sa.cell(row=r, column=4).value
        action = "删除" if vb in (None, "") else "改写"
        print(f"[{action}] 序号{label} · {where} · {element} · {scene}")
        print(f"        原：{va}")
        print(f"        新：{vb}")
print(f"total {total}")
