"""Render a sheet of the workbook to HTML so it can be eyeballed without Excel."""

import html
import os
import subprocess
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "记账本文案.xlsx")
OUT = os.path.join(HERE, "_preview.html")
PNG = os.path.join(HERE, "_preview.png")
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

sheet = sys.argv[1] if len(sys.argv) > 1 else "界面文字"
rows_wanted = int(sys.argv[2]) if len(sys.argv) > 2 else 36

book = openpyxl.load_workbook(XLSX)
page = book[sheet]
head = page.max_row and [c.value for c in page[3]]
body = list(page.iter_rows(min_row=4, max_row=3 + rows_wanted, values_only=True))

widths = {"界面文字": [40, 140, 160, 100, 420, 420],
          "导出文件里的文字": [40, 150, 150, 100, 420, 420],
          "未列入": [230, 50, 330, 380]}.get(sheet, [200] * 6)

parts = []
if head and head[0] is not None:
    parts.append("<tr>" + "".join(
        f"<td>{html.escape(str(c)) if c is not None else ''}</td>" for c in head) + "</tr>")
for record in body:
    parts.append("<tr>" + "".join(
        f"<td>{html.escape(str(c)) if c is not None else ''}</td>" for c in record) + "</tr>")

style = "".join(
    f"td:nth-child({index + 1}){{width:{width}px}}"
    for index, width in enumerate(widths))
doc = (
    '<!doctype html><meta charset="utf-8">'
    '<style>body{margin:0;background:#fff;font:13px/1.45 system-ui,"Microsoft YaHei"}'
    'table{border-collapse:collapse}'
    'td{border:1px solid #d9d9d9;padding:4px 6px;vertical-align:top}'
    'tr:first-child td{background:#1f3864;color:#fff;font-weight:700}'
    'tr:nth-child(even) td{background:#f2f5fa}'
    f'td:last-child{{background:#fff2cc}}{style}</style><table>'
    + "".join(parts) + "</table>")
open(OUT, "w", encoding="utf-8").write(doc)

subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars",
                "--window-size=1400,1180", f"--screenshot={PNG}",
                "file:///" + OUT.replace("\\", "/")], check=True, capture_output=True)
print("wrote", PNG)
